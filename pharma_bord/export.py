"""Exports : bon de commande (HTML imprimable et CSV), inventaire (CSV, Excel), historique (CSV)."""

from __future__ import annotations

import csv
import html
from datetime import date
from pathlib import Path
from typing import Dict, Iterable, List, Optional

from .db import ConfigNavire, Mouvement
from .moteur import Analyse, EtatArticle, LIBELLES_DOTATION
from .referentiel import TYPE_MED


def _fmt_date(d) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def _txt(v) -> str:
    return "" if v is None else str(v)


# ---------------------------------------------------------------------- commande
class LigneCommande:
    """Un article à commander : état + quantité demandée (calculée ou ajustée à la main)."""

    __slots__ = ("etat", "quantite", "ajustee")

    def __init__(self, etat: EtatArticle, quantite: int, ajustee: bool):
        self.etat, self.quantite, self.ajustee = etat, quantite, ajustee


def lignes_commande(analyse: Analyse, ajustements: Optional[Dict[str, int]] = None,
                    inclure_tout: bool = False) -> List[LigneCommande]:
    """Articles de la commande. ``ajustements`` : item_id -> quantité demandée imposée (0 = retirer)."""
    aj = ajustements or {}
    res = []
    for e in analyse.etats:
        if e.exigence is None and e.article.item_id not in aj:
            continue
        if e.article.item_id in aj:
            q, ajustee = aj[e.article.item_id], True
        else:
            q, ajustee = e.a_commander, False
        if q > 0 or inclure_tout:
            res.append(LigneCommande(e, q, ajustee))
    return res


def commande_csv(analyse: Analyse, chemin: Path, ajustements: Optional[Dict[str, int]] = None) -> None:
    with open(chemin, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["Type", "Catégorie", "ATC", "Dénomination", "Voie", "Forme / présentation", "Liste",
                    "Quantité réglementaire", "Stock valide à l'horizon", "Quantité demandée", "Ajustée à la main",
                    "Prochaine péremption", "Remarques"])
        for lc in lignes_commande(analyse, ajustements):
            e, a = lc.etat, lc.etat.article
            w.writerow(["Médicament" if a.type == TYPE_MED else "Matériel", e.categorie, a.atc,
                        a.libelle if a.kit else a.nom, a.voie, a.forme, a.liste, _txt(e.requis),
                        e.stock_valide_horizon, lc.quantite, "oui" if lc.ajustee else "",
                        _fmt_date(e.prochaine_peremption), a.remarques])


TITRES_OFFICIELS = {
    "A": ("Dotation médicale A", ""),
    "B": ("Dotation médicale B", ""),
    "BR": ("Dotation médicale B restreinte",
           "Pour les navires en deuxième catégorie de navigation restreinte et ne s'éloignant pas à plus de 60 milles marins des côtes"),
    "C": ("Dotation médicale C", ""),
    "CR": ("Dotation médicale C restreinte",
           "Pour les navires ne s'éloignant pas à plus de 20 milles marins des côtes et restant moins de 24 heures à la mer"),
    "CAT5": ("Dotation médicale 5e catégorie", ""),
}

_STYLE_BON = """
 @page { size: A4; margin: 12mm; }
 body { font-family: Arial, Helvetica, sans-serif; font-size: 10.5px; margin: 16px; color: #000; }
 .aide { background: #fff8d6; border: 1px solid #e0c060; padding: 6px 10px; margin-bottom: 12px; font-size: 11px; }
 .aide button { font-size: 12px; padding: 4px 12px; margin-right: 10px; }
 h1 { font-size: 17px; text-align: center; margin: 4px 0 0; }
 .sous { text-align: center; font-size: 11px; margin: 0 0 2px; }
 h2 { font-size: 13px; text-align: center; margin: 2px 0; }
 .service { text-align: center; font-size: 11px; margin-bottom: 10px; }
 table.entete { width: 100%; border-collapse: collapse; margin-bottom: 10px; }
 table.entete td { border: 1px solid #000; padding: 6px 8px; vertical-align: top; width: 50%; }
 table.entete .champ { display: inline-block; min-width: 60%; border-bottom: 1px dotted #666; padding: 0 3px; }
 .tampon { border: 1px dashed #666; height: 70px; margin-top: 6px; text-align: center; color: #666; line-height: 70px; }
 table.lignes { width: 100%; border-collapse: collapse; }
 table.lignes th, table.lignes td { border: 1px solid #000; padding: 3px 4px; vertical-align: top; }
 table.lignes th { background: #e8e8e8; font-size: 10px; }
 table.lignes td.num { text-align: center; white-space: nowrap; }
 table.lignes tr.cat td { background: #f2f2f2; font-weight: bold; }
 table.lignes td.aj { background: #eef6ff; }
 .stup { font-weight: bold; }
 [contenteditable]:hover { outline: 1px dashed #4a90d9; }
 table.sign { width: 100%; border-collapse: collapse; margin-top: 16px; page-break-inside: avoid; }
 table.sign td { border: 1px solid #000; padding: 6px 8px; vertical-align: top; width: 33%; height: 95px; }
 .section { page-break-before: always; }
 .note { font-size: 9.5px; margin-top: 10px; }
 @media print { .aide { display: none; } body { margin: 0; } [contenteditable]:hover { outline: none; } }
"""


def commande_html(analyse: Analyse, cfg: ConfigNavire, chemin: Path, inclure_tout: bool = False,
                  ajustements: Optional[Dict[str, int]] = None) -> None:
    """Bon de commande au format officiel de l'annexe 217-3.A.5 (constitution / renouvellement).

    Les cases (en-tête, spécialité délivrée, quantités) sont modifiables dans le navigateur avant
    impression ; le cadre « tampon humide » est réservé au pharmacien.
    """
    lignes = lignes_commande(analyse, ajustements, inclure_tout)
    esc = html.escape
    titre, sous_titre = TITRES_OFFICIELS.get(cfg.dotation_base, (f"Dotation médicale {cfg.dotation_base}", ""))
    if cfg.complement_passagers:
        titre += f" + complément passagers {cfg.complement_passagers}"
    ed = ' contenteditable="true"'
    navire = esc(cfg.nom) + (f" – {esc(cfg.immatriculation)}" if cfg.immatriculation else "")
    parts = [f"""<!DOCTYPE html>
<html lang="fr"><head><meta charset="utf-8">
<title>{esc(titre)} – Bon de commande</title>
<style>{_STYLE_BON}</style></head><body>
<div class="aide"><button onclick="window.print()">Imprimer</button>
Toutes les cases (en-tête, spécialité délivrée, quantités demandées) sont modifiables : cliquez dans une case pour
corriger le texte avant impression. Le cadre « Tampon humide » est réservé à la pharmacie.</div>
"""]

    def entete(sous_type: str) -> str:
        effectif = f"{cfg.effectif} marins" + (f" – {cfg.nb_passagers} passagers" if cfg.nb_passagers else "")
        sous = f'<p class="sous">{esc(sous_titre)}</p>' if sous_titre else ""
        return f"""
<h1>{esc(titre)}</h1>{sous}
<h2>Bon de commande pour constitution / renouvellement – {esc(sous_type)}</h2>
<div class="service">Service de santé des gens de mer – Direction des affaires maritimes</div>
<table class="entete"><tr>
<td><b>Armement :</b> <span class="champ"{ed}>{esc(cfg.armement)}</span><br><br>
<b>Navire :</b> <span class="champ"{ed}>{navire}</span><br><br>
<b>Nom du Capitaine / Armateur :</b> <span class="champ"{ed}>{esc(cfg.capitaine)}</span><br><br>
<b>Adresse :</b> <span class="champ"{ed}>{esc(cfg.adresse_armement)}</span><br><br>
<b>Effectif :</b> <span class="champ"{ed}>{esc(effectif)}</span><br><br>
<b>Date de la commande :</b> <span class="champ"{ed}>{_fmt_date(analyse.aujourdhui)}</span></td>
<td><b>Établissement pharmaceutique / Pharmacien</b><br><br>
<b>Nom :</b> <span class="champ"{ed}>{esc(cfg.pharmacien)}</span><br><br>
<b>Adresse :</b> <span class="champ"{ed}>{esc(cfg.adresse_pharmacien)}</span>
<div class="tampon">(Tampon humide)</div></td>
</tr></table>
"""

    def signatures() -> str:
        return f"""
<table class="sign"><tr>
<td><b>Commande établie par le capitaine ou le responsable des soins à bord</b><br><br>
Nom : <span{ed}>{esc(cfg.responsable_soins or cfg.capitaine)}</span><br>
Fonction : <span{ed}>{esc(cfg.fonction_responsable)}</span><br><br>Signature :</td>
<td><b>Procuration</b><br><br>Le Capitaine / Armement donne procuration à : M./Mme <span{ed}></span><br>
(nom – prénom)<br><br>Signature :</td>
<td><b>Signature</b><br><br><br>Le Capitaine / Armement</td>
</tr></table>
<p class="note">Sauf cas de force majeure, les médicaments relevant de la réglementation des stupéfiants, de la liste I ou de
la liste II ne peuvent être utilisés que sur prescription médicale. Les stupéfiants sont délivrés sur présentation du
registre des stupéfiants et de l'état récapitulatif d'administration (art. 217-4.03). Les médicaments à usage
hospitalier sont délivrés conformément à l'art. L. 5126-5 du code de la santé publique.</p>
"""

    meds = [lc for lc in lignes if lc.etat.article.type == TYPE_MED]
    mats = [lc for lc in lignes if lc.etat.article.type != TYPE_MED]

    # --- médicaments : format officiel
    parts.append(entete("Médicaments"))
    parts.append("<table class='lignes'><thead><tr><th>Numéro ATC</th><th>Dénomination commune internationale</th>"
                 "<th>Nom de la spécialité délivrée</th><th>Forme – composition</th><th>Quantité réglementaire</th>"
                 "<th>Quantité demandée</th><th>Quantité délivrée</th><th>Liste</th></tr></thead><tbody>")
    if not meds:
        parts.append("<tr><td colspan='8' style='text-align:center'>Aucun médicament à commander.</td></tr>")
    cat = None
    for lc in meds:
        e, a = lc.etat, lc.etat.article
        if e.categorie != cat:
            cat = e.categorie
            parts.append(f"<tr class='cat'><td colspan='8'>{esc(cat)}</td></tr>")
        liste = f"<span class='stup'>{esc(a.liste)}</span>" if a.est_stupefiant else esc(a.liste)
        regl = esc(_txt(e.requis)) if e.requis is not None else esc(e.exigence.quantite_texte if e.exigence else "")
        voie = f" <i>({esc(a.voie)})</i>" if a.voie and a.voie != "orale" else ""
        parts.append(f"<tr><td>{esc(a.atc)}</td><td>{esc(a.nom)}{voie}</td>"
                     f"<td{ed}></td><td>{esc(a.forme)}</td><td class='num'>{regl}</td>"
                     f"<td class='num{' aj' if lc.ajustee else ''}'{ed}><b>{lc.quantite or ''}</b></td>"
                     f"<td class='num'{ed}></td><td class='num'>{liste}</td></tr>")
    parts.append("</tbody></table>")
    parts.append(signatures())

    # --- matériel : même présentation, page séparée
    if mats:
        parts.append("<div class='section'></div>")
        parts.append(entete("Matériel médical et objets de pansement"))
        parts.append("<table class='lignes'><thead><tr><th>Article</th><th>Présentation</th><th>Remarques / type</th>"
                     "<th>Référence délivrée</th><th>Quantité réglementaire</th><th>Quantité demandée</th>"
                     "<th>Quantité délivrée</th></tr></thead><tbody>")
        cat = None
        for lc in mats:
            e, a = lc.etat, lc.etat.article
            if e.categorie != cat:
                cat = e.categorie
                parts.append(f"<tr class='cat'><td colspan='7'>{esc(cat)}</td></tr>")
            regl = esc(_txt(e.requis)) if e.requis is not None else esc(e.exigence.quantite_texte if e.exigence else "")
            nom = ("[" + a.kit.capitalize() + "] " if a.kit else "") + a.nom
            parts.append(f"<tr><td>{esc(nom)}</td><td>{esc(a.forme)}</td><td>{esc(a.remarques)}</td><td{ed}></td>"
                         f"<td class='num'>{regl}</td><td class='num{' aj' if lc.ajustee else ''}'{ed}><b>{lc.quantite or ''}</b></td>"
                         f"<td class='num'{ed}></td></tr>")
        parts.append("</tbody></table>")
        parts.append(signatures())
    parts.append("</body></html>")
    Path(chemin).write_text("".join(parts), encoding="utf-8")


# ---------------------------------------------------------------------- inventaire
def inventaire_csv(analyse: Analyse, chemin: Path) -> None:
    """Inventaire de la dotation médicale : une boîte par ligne."""
    with open(chemin, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["Type", "Catégorie", "Dénomination", "Voie", "Forme / présentation", "Liste",
                    "Quantité réglementaire", "Statut", "Lot", "Spécialité", "Quantité du lot",
                    "Péremption", "Emplacement", "Entrée le", "Commentaire"])
        for e in analyse.etats:
            a = e.article
            base = ["Médicament" if a.type == TYPE_MED else "Matériel", e.categorie,
                    ("[" + a.kit.capitalize() + "] " if a.kit else "") + a.nom, a.voie, a.forme, a.liste,
                    _txt(e.requis) if e.requis is not None else (e.exigence.quantite_texte if e.exigence else ""),
                    e.statut]
            if not e.lots:
                w.writerow(base + ["", "", 0, "", "", "", ""])
            for l in e.lots:
                w.writerow(base + [l.numero_lot, l.specialite, l.quantite, _fmt_date(l.peremption),
                                   l.emplacement, _fmt_date(l.date_entree), l.commentaire])


def inventaire_xlsx(analyse: Analyse, cfg: ConfigNavire, chemin: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventaire"
    ws.append([f"Inventaire de la dotation médicale – {cfg.nom} ({cfg.immatriculation})",
               f"Dotation {cfg.dotation_base}{(' + ' + cfg.complement_passagers) if cfg.complement_passagers else ''}",
               f"Édité le {_fmt_date(analyse.aujourdhui)}", f"Horizon {_fmt_date(analyse.horizon)}"])
    hdr = ["Type", "Catégorie", "Dénomination", "Voie", "Forme / présentation", "Liste", "Qté réglementaire",
           "Stock total", "Stock valide à l'horizon", "À commander", "Statut", "Lot", "Spécialité",
           "Qté du lot", "Péremption", "Emplacement", "Entrée le", "Commentaire"]
    ws.append(hdr)
    for i in range(1, len(hdr) + 1):
        c = ws.cell(row=2, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for w, i in zip([11, 30, 45, 16, 40, 10, 10, 10, 12, 10, 18, 14, 22, 9, 12, 16, 12, 30], range(1, 19)):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    for e in analyse.etats:
        a = e.article
        base = ["Médicament" if a.type == TYPE_MED else "Matériel", e.categorie,
                ("[" + a.kit.capitalize() + "] " if a.kit else "") + a.nom, a.voie, a.forme, a.liste,
                e.requis if e.requis is not None else (e.exigence.quantite_texte if e.exigence else ""),
                e.stock_total, e.stock_valide_horizon, e.a_commander, e.statut]
        if not e.lots:
            ws.append(base + ["", "", 0, "", "", "", ""])
        for l in e.lots:
            ws.append(base + [l.numero_lot, l.specialite, l.quantite, l.peremption, l.emplacement,
                              l.date_entree, l.commentaire])
    for row in ws.iter_rows(min_row=3):
        for c in (row[14], row[16]):
            c.number_format = "DD/MM/YYYY"
    wb.save(chemin)


def mouvements_csv(mvts: Iterable[Mouvement], libelles: dict, chemin: Path) -> None:
    """Historique des mouvements de l'inventaire (ne remplace pas les registres officiels du bord)."""
    with open(chemin, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["Date", "Article", "Boîte", "Type", "Quantité", "Motif", "Commentaire"])
        for m in mvts:
            w.writerow([m.horodatage.strftime("%d/%m/%Y %H:%M"), libelles.get(m.item_id, m.item_id),
                        m.lot_id or "", m.type, m.quantite, m.motif, m.commentaire])
